# app.py - Professional DocuMind AI for Hugging Face Spaces
import gradio as gr
import os
import tempfile
from typing import List, Dict, Any
import fitz  # PyMuPDF
from PIL import Image
import pytesseract
import docx
import re
from datetime import datetime

# Core imports
import chromadb
from sentence_transformers import SentenceTransformer
import google.generativeai as genai
from langchain.text_splitter import RecursiveCharacterTextSplitter

# Configure Gemini (will use environment variable)
if os.getenv("GEMINI_API_KEY"):
    genai.configure(api_key=os.getenv("GEMINI_API_KEY"))

class DocumentProcessor:
    """Professional document processing for multiple file types."""
    
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
    
    def process_file(self, file_path: str, file_name: str) -> tuple:
        """Process uploaded file and return text content and metadata."""
        try:
            file_extension = os.path.splitext(file_name)[1].lower()
            
            if file_extension == '.pdf':
                text = self._process_pdf(file_path)
            elif file_extension == '.docx':
                text = self._process_docx(file_path)
            elif file_extension in ['.png', '.jpg', '.jpeg']:
                text = self._process_image(file_path)
            elif file_extension == '.txt':
                text = self._process_text(file_path)
            else:
                return None, f"Unsupported file type: {file_extension}"
            
            if not text.strip():
                return None, "No text could be extracted from the file"
            
            # Create chunks
            chunks = self.create_chunks(text)
            
            metadata = {
                "filename": file_name,
                "file_type": file_extension,
                "text_length": len(text),
                "chunks_created": len(chunks),
                "processed_at": datetime.now().isoformat()
            }
            
            return {"text": text, "chunks": chunks}, metadata
            
        except Exception as e:
            return None, f"Error processing file: {str(e)}"
    
    def _process_pdf(self, file_path: str) -> str:
        """Extract text from PDF."""
        text_content = ""
        doc = fitz.open(file_path)
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text.strip():
                text_content += f"\n--- Page {page_num + 1} ---\n{text}"
        
        doc.close()
        return self._clean_text(text_content)
    
    def _process_docx(self, file_path: str) -> str:
        """Extract text from Word document."""
        doc = docx.Document(file_path)
        text_content = ""
        
        for paragraph in doc.paragraphs:
            text_content += paragraph.text + "\n"
        
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text for cell in row.cells])
                text_content += row_text + "\n"
        
        return self._clean_text(text_content)
    
    def _process_image(self, file_path: str) -> str:
        """Extract text from image using OCR with better error handling."""
        try:
            from PIL import Image
            import pytesseract
            
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            
            if not text.strip():
                return "No text detected in this image. Please try an image with clearer text or use a text-based document."
            
            return self._clean_text(text)
            
        except Exception as e:
            error_msg = str(e).lower()
            if "tesseract" in error_msg or "not installed" in error_msg:
                return """OCR functionality is currently unavailable. 
                
                For best results, please upload:
                - PDF files with text
                - Word documents (.docx) 
                - Plain text files (.txt)
                
                Image OCR will be available in the next update."""
            else:
                return f"Unable to process this image: {str(e)}"
    
    def _process_text(self, file_path: str) -> str:
        """Process plain text file."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return self._clean_text(content)
    
    def _clean_text(self, text: str) -> str:
        """Clean and normalize text."""
        text = re.sub(r'\s+', ' ', text)
        text = re.sub(r'[^\w\s\-.,;:!?()]', '', text)
        return text.strip()
    
    def create_chunks(self, text: str) -> List[Dict[str, Any]]:
        """Split text into chunks."""
        if not text.strip():
            return []
        
        chunks = self.text_splitter.split_text(text)
        
        chunk_data = []
        for i, chunk in enumerate(chunks):
            chunk_data.append({
                "content": chunk,
                "chunk_index": i,
                "word_count": len(chunk.split()),
                "char_count": len(chunk)
            })
        
        return chunk_data

class VectorStore:
    """Professional vector store for semantic search."""
    
    def __init__(self):
        self.client = chromadb.Client()
        self.embedding_model = SentenceTransformer('all-MiniLM-L6-v2')
        self.collection = None
        self._create_collection()
    
    def _create_collection(self):
        """Create or recreate the collection."""
        try:
            if self.collection:
                self.client.delete_collection("documents")
        except:
            pass
        
        self.collection = self.client.create_collection(
            name="documents",
            metadata={"hnsw:space": "cosine"}
        )
    
    def add_document(self, chunks: List[Dict[str, Any]], document_name: str) -> bool:
        """Add document chunks to vector store."""
        if not chunks:
            return False
        
        try:
            ids = []
            documents = []
            metadatas = []
            embeddings = []
            
            for chunk in chunks:
                chunk_id = f"{document_name}_{chunk['chunk_index']}"
                embedding = self.embedding_model.encode(chunk['content']).tolist()
                
                metadata = {
                    "document_name": document_name,
                    "chunk_index": chunk['chunk_index'],
                    "word_count": chunk['word_count'],
                    "char_count": chunk['char_count']
                }
                
                ids.append(chunk_id)
                documents.append(chunk['content'])
                metadatas.append(metadata)
                embeddings.append(embedding)
            
            self.collection.add(
                ids=ids,
                documents=documents,
                metadatas=metadatas,
                embeddings=embeddings
            )
            
            return True
            
        except Exception as e:
            print(f"Error adding to vector store: {e}")
            return False
    
    def search_similar(self, query: str, n_results: int = 3) -> List[Dict[str, Any]]:
        """Search for similar chunks."""
        if not query.strip():
            return []
        
        try:
            if self.collection.count() == 0:
                return []
            
            query_embedding = self.embedding_model.encode(query).tolist()
            
            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=min(n_results, self.collection.count()),
                include=["documents", "metadatas", "distances"]
            )
            
            formatted_results = []
            if results['documents'] and results['documents'][0]:
                for i in range(len(results['documents'][0])):
                    distance = results['distances'][0][i]
                    similarity_score = max(0, 1 - distance)
                    
                    result = {
                        "content": results['documents'][0][i],
                        "metadata": results['metadatas'][0][i],
                        "similarity_score": round(similarity_score, 3),
                        "chunk_id": results['ids'][0][i]
                    }
                    formatted_results.append(result)
            
            return formatted_results
            
        except Exception as e:
            print(f"Error searching: {e}")
            return []

class LLMService:
    """Professional LLM service with Gemini."""
    
    def __init__(self):
        self.client = None
        if os.getenv("GEMINI_API_KEY"):
            try:
                self.client = genai.GenerativeModel('models/gemini-2.0-flash')
            except:
                pass
    
    def generate_answer(self, question: str, context_chunks: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Generate answer using Gemini."""
        if not self.client:
            return {
                "answer": "⚠️ Gemini API not configured. Please set GEMINI_API_KEY environment variable for AI responses.",
                "confidence": 0.0,
                "sources_used": []
            }
        
        try:
            context = self._format_context(context_chunks)
            
            prompt = f"""You are DocuMind AI, a professional document analysis assistant.

INSTRUCTIONS:
- Answer based ONLY on the provided document context
- If information isn't available in context, state that clearly
- Be precise and cite relevant sources
- Maintain a professional tone

DOCUMENT CONTEXT:
{context}

USER QUESTION: {question}

Please provide a clear, accurate answer based on the context provided."""

            response = self.client.generate_content(prompt)
            
            if response and response.text:
                return {
                    "answer": response.text,
                    "confidence": self._estimate_confidence(context_chunks),
                    "sources_used": self._extract_sources(context_chunks)
                }
            
        except Exception as e:
            return {
                "answer": f"Error generating response: {str(e)}",
                "confidence": 0.0,
                "sources_used": []
            }
        
        return {
            "answer": "Unable to generate response. Please try again.",
            "confidence": 0.0,
            "sources_used": []
        }
    
    def _format_context(self, context_chunks: List[Dict[str, Any]]) -> str:
        """Format context for the prompt."""
        if not context_chunks:
            return "No relevant context found."
        
        formatted_context = []
        for i, chunk in enumerate(context_chunks[:3], 1):
            metadata = chunk.get('metadata', {})
            source_name = metadata.get('document_name', 'Unknown')
            similarity = chunk.get('similarity_score', 0)
            
            formatted_context.append(
                f"[Source {i}: {source_name} | Relevance: {similarity:.3f}]\n"
                f"{chunk['content']}\n"
            )
        
        return "\n".join(formatted_context)
    
    def _extract_sources(self, context_chunks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Extract source information."""
        sources = []
        seen_docs = set()
        
        for chunk in context_chunks:
            metadata = chunk.get('metadata', {})
            doc_name = metadata.get('document_name', 'Unknown')
            
            if doc_name not in seen_docs:
                sources.append({
                    "document_name": doc_name,
                    "similarity_score": chunk.get('similarity_score', 0),
                    "chunk_preview": chunk['content'][:100] + "..." if len(chunk['content']) > 100 else chunk['content']
                })
                seen_docs.add(doc_name)
        
        return sources
    
    def _estimate_confidence(self, context_chunks: List[Dict[str, Any]]) -> float:
        """Estimate confidence based on context quality."""
        if not context_chunks:
            return 0.0
        
        scores = [chunk.get('similarity_score', 0) for chunk in context_chunks]
        avg_score = sum(scores) / len(scores)
        return round(min(avg_score * 1.1, 1.0), 2)

class DocuMindAI:
    """Main DocuMind AI system."""
    
    def __init__(self):
        self.processor = DocumentProcessor()
        self.vector_store = VectorStore()
        self.llm_service = LLMService()
        self.documents = {}
    
    def process_document(self, file_path: str, file_name: str) -> tuple:
        """Process and add document to the system."""
        document_data, metadata = self.processor.process_file(file_path, file_name)
        
        if document_data is None:
            return False, metadata
        
        success = self.vector_store.add_document(document_data["chunks"], file_name)
        
        if success:
            self.documents[file_name] = metadata
            return True, f"✅ Successfully processed '{file_name}' - {metadata['chunks_created']} chunks created"
        else:
            return False, "❌ Failed to add document to vector store"
    
    def ask_question(self, question: str) -> Dict[str, Any]:
        """Ask a question about the documents."""
        if not question.strip():
            return {
                "answer": "Please enter a question.",
                "confidence": 0.0,
                "sources": []
            }
        
        if not self.documents:
            return {
                "answer": "Please upload documents first before asking questions.",
                "confidence": 0.0,
                "sources": []
            }
        
        # Search for relevant context
        context_chunks = self.vector_store.search_similar(question, n_results=3)
        
        if not context_chunks:
            return {
                "answer": "I couldn't find relevant information in your documents to answer this question.",
                "confidence": 0.0,
                "sources": []
            }
        
        # Generate answer
        result = self.llm_service.generate_answer(question, context_chunks)
        
        return {
            "answer": result["answer"],
            "confidence": result["confidence"],
            "sources": result.get("sources_used", [])
        }
    
    def get_document_stats(self) -> str:
        """Get statistics about loaded documents."""
        if not self.documents:
            return "No documents loaded."
        
        stats = []
        total_chunks = 0
        
        for doc_name, metadata in self.documents.items():
            stats.append(f"📄 **{doc_name}** ({metadata['file_type']}) - {metadata['chunks_created']} chunks")
            total_chunks += metadata['chunks_created']
        
        return f"**📊 Document Statistics:**\n\n" + "\n".join(stats) + f"\n\n**Total chunks in database: {total_chunks}**"

# Initialize the system
documind_ai = DocuMindAI()

def process_file_upload(file):
    """Handle file upload and processing."""
    if file is None:
        return "Please upload a file first.", "", ""
    
    try:
        # Process the uploaded file
        success, message = documind_ai.process_document(file.name, os.path.basename(file.name))
        
        if success:
            stats = documind_ai.get_document_stats()
            return message, stats, "✅ Document ready! You can now ask questions about it."
        else:
            return f"❌ Error: {message}", "", ""
    
    except Exception as e:
        return f"❌ Error processing file: {str(e)}", "", ""

def answer_question(question, chat_history):
    """Handle question answering."""
    if not question.strip():
        return "", chat_history
    
    # Get answer from DocuMind AI
    result = documind_ai.ask_question(question)
    
    # Format the response
    response = f"**🤖 DocuMind AI:**\n{result['answer']}\n\n"
    
    if result['sources']:
        response += f"**📚 Sources ({len(result['sources'])}):**\n"
        for i, source in enumerate(result['sources'], 1):
            response += f"{i}. {source['document_name']} (relevance: {source['similarity_score']:.3f})\n"
        #response += f"\n**📊 Confidence: {result['confidence']:.2f}**"
    
    # Update chat history
    chat_history.append([question, response])
    
    return "", chat_history

def clear_documents():
    """Clear all documents and restart system."""
    global documind_ai
    documind_ai = DocuMindAI()
    return "🔄 All documents cleared. Upload new documents to start fresh.", "", ""

# Create the Gradio interface
with gr.Blocks(
    title="DocuMind AI - Professional Document Analysis",
    theme=gr.themes.Soft(),
    css="""
    .gradio-container {
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    .header {
        text-align: center;
        background: linear-gradient(90deg, #667eea 0%, #764ba2 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-size: 2.5rem;
        font-weight: bold;
        margin-bottom: 1rem;
    }
    """
) as interface:
    
    # Header
    gr.HTML("""
    <div class="header">
        🧠 DocuMind AI
    </div>
    <p style="text-align: center; font-size: 1.2rem; color: #666; margin-bottom: 2rem;">
        Professional Document Analysis with AI-Powered Question Answering
    </p>
    """)
    
    with gr.Row():
        with gr.Column(scale=1):
            gr.Markdown("## 📁 Document Upload")
            
            file_input = gr.File(
                label="Upload Document",
                file_types=[".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg"],
                file_count="single"
            )
            
            upload_btn = gr.Button("🚀 Process Document", variant="primary", size="lg")
            
            upload_status = gr.Textbox(
                label="Processing Status",
                interactive=False,
                max_lines=3
            )
            
            doc_stats = gr.Markdown("Upload a document to see statistics here.")
            
            clear_btn = gr.Button("🗑️ Clear All Documents", variant="secondary")
        
        with gr.Column(scale=2):
            gr.Markdown("## 💬 Ask Questions")
            
            chatbot = gr.Chatbot(
                label="DocuMind AI Assistant",
                height=400,
                show_copy_button=True
            )
            
            with gr.Row():
                question_input = gr.Textbox(
                    label="Your Question",
                    placeholder="Ask anything about your uploaded documents...",
                    scale=4
                )
                ask_btn = gr.Button("📤 Ask", variant="primary", scale=1)
            
            chat_status = gr.Textbox(
                label="Chat Status",
                interactive=False,
                visible=False
            )
    
    # Example questions
    gr.Markdown("""
    ### 💡 Example Questions:
    - "What is this document about?"
    - "Summarize the key points"
    - "What are the main features mentioned?"
    - "Who are the target users?"
    - "What technical details are provided?"
    """)
    
    # Event handlers
    upload_btn.click(
        fn=process_file_upload,
        inputs=[file_input],
        outputs=[upload_status, doc_stats, chat_status]
    )
    
    ask_btn.click(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[question_input, chatbot]
    )
    
    question_input.submit(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[question_input, chatbot]
    )
    
    clear_btn.click(
        fn=clear_documents,
        outputs=[upload_status, doc_stats, chat_status]
    )

# Launch the interface
if __name__ == "__main__":
    interface.launch(
        server_name="0.0.0.0",
        server_port=7860,
        show_error=True
    )